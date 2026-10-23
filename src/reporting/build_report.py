"""
Render the final report from Markdown source to .docx (and optionally PDF via Microsoft Word).

Supported Markdown subset: `#`/`##`/`###` headings, paragraphs, `-` and `1.` lists, pipe tables,
images `![Figure N. caption](path)`, table captions as a paragraph starting with `Table N.`,
**bold**, *italic*, `code`, and a `<!-- pagebreak -->` marker.

Formatting: Times New Roman 12 pt body, 1.5 line spacing, 1-inch margins, page numbers in the footer.

Usage:
    python -m src.reporting.build_report reports/final_report.md --pdf
"""

from __future__ import annotations

import argparse
import re
import subprocess
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from src.config import REPO_ROOT

FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _set_font(run, size: float | None = None, bold=None, italic=None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(par, text: str, size: float | None = None, base_bold=False, base_italic=False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _set_font(par.add_run(part[2:-2]), size, bold=True, italic=base_italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _set_font(par.add_run(part[1:-1]), size, bold=base_bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            _set_font(r, size, bold=base_bold, italic=base_italic)
            r.font.name = "Consolas"
        else:
            _set_font(par.add_run(part), size, bold=base_bold, italic=base_italic)


def _page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    _set_font(run, 11)


def _style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    for lvl, size in ((1, 14), (2, 13), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = None
        st.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
        _page_number_footer(s)


def _table(doc: Document, rows: List[List[str]]) -> None:
    header, body = rows[0], [r for r in rows[1:] if not re.fullmatch(r"[\s:|-]+", "|".join(r))]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + body):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, cell.strip(), size=9.5, base_bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(md_path: Path, out_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _style_document(doc)
    base = md_path.parent
    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("<!--"):
            while "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1 and not title_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text, size=16, base_bold=True)
                title_done = True
            else:
                h = doc.add_heading(level=max(1, level - 1))
                add_inline(h, text, size={1: 14, 2: 13, 3: 12}[max(1, level - 1)], base_bold=True)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(doc, rows)
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", line)
        if m:
            caption, src = m.group(1), m.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((base / src).resolve()), width=Inches(6.2))
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_inline(c, caption, size=10, base_italic=False)
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_inline(p, m.group(3))
            i += 1
            continue
        # Paragraph: merge consecutive non-special lines
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\s*([-*]|\d+\.)\s|<!--)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)
        p = doc.add_paragraph()
        if re.match(r"^\*\*Table \d+\.", text) or re.match(r"^Table \d+\.", text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_inline(p, text, size=10)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, text)
    doc.save(out_path)
    return out_path


def to_pdf(docx_path: Path) -> Path:
    """Export with Microsoft Word through COM (Windows); returns the PDF path."""
    pdf = docx_path.with_suffix(".pdf")
    script = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
              f"$d = $w.Documents.Open('{docx_path}'); $d.Fields.Update() | Out-Null; "
              f"$d.SaveAs([ref]'{pdf}', [ref]17); $d.Close(); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return pdf


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=str)
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    md = (REPO_ROOT / args.markdown).resolve()
    out = render(md, md.with_suffix(".docx"))
    print("Wrote", out)
    if args.pdf:
        print("Wrote", to_pdf(out))


if __name__ == "__main__":
    main()
