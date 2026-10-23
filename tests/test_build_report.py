"""
Tests for the report renderer (`src.reporting.build_report`): formatting requirements.
"""

import tempfile
import unittest
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_LINE_SPACING  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

from src.reporting.build_report import render  # noqa: E402


class TestRender(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        fig, ax = plt.subplots(figsize=(2, 1))
        ax.plot([0, 1], [0, 1])
        fig.savefig(self.dir / "f.png")
        plt.close(fig)
        (self.dir / "r.md").write_text(
            "# A Title\n\n## 1. Introduction\n\nBody with **bold** and *italic* text.\n\n"
            "- point one\n- point two\n\n**Table 1.** Caption.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
            "![Figure 1. A figure.](f.png)\n\n<!-- pagebreak -->\n\n## References\n\nRef.\n",
            encoding="utf-8")
        self.doc = Document(str(render(self.dir / "r.md", self.dir / "r.docx")))

    def test_body_formatting(self):
        normal = self.doc.styles["Normal"]
        self.assertEqual(normal.font.name, "Times New Roman")
        self.assertEqual(normal.font.size, Pt(12))
        self.assertEqual(normal.paragraph_format.line_spacing_rule, WD_LINE_SPACING.ONE_POINT_FIVE)

    def test_margins_and_page_numbers(self):
        s = self.doc.sections[0]
        for m in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin):
            self.assertEqual(m, Inches(1))
        self.assertIn("PAGE", s.footer.paragraphs[0]._p.xml)

    def test_content_elements(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertEqual(texts[0], "A Title")
        self.assertIn("1. Introduction", texts)
        self.assertEqual(len(self.doc.tables), 1)
        self.assertEqual(self.doc.tables[0].cell(1, 1).text, "2")
        self.assertEqual(len(self.doc.inline_shapes), 1)
        self.assertIn("Figure 1. A figure.", texts)
        runs = [r for p in self.doc.paragraphs for r in p.runs if r.text == "bold"]
        self.assertTrue(runs and runs[0].bold)


class TestReportTables(unittest.TestCase):
    """Result tables are generated from saved artifacts; empty placeholders and NaN must be handled."""

    def setUp(self):
        import pandas as pd
        self.tables = Path(tempfile.mkdtemp())
        pd.DataFrame([
            {"series_id": "potato_uttar_pradesh", "method": "robust", "evaluated_days": 10, "label_base_rate": 0.5,
             "tp": 0, "fp": 0, "fn": 5, "precision": float("nan"), "recall": 0.0, "f1": 0.0,
             "null_precision_median": float("nan"), "randomisation_p_precision": float("nan")},
        ]).to_csv(self.tables / "anomaly_weak_label_agreement.csv", index=False)

    def test_empty_placeholder_is_filled_and_nan_shown_as_dash(self):
        from src.reporting.report_tables import refresh
        text = "a\n<!-- BEGIN TABLE:anomaly_weak -->\n<!-- END TABLE:anomaly_weak -->\nb\n"
        out = refresh(text, self.tables)
        row = [l for l in out.splitlines() if l.startswith("| Potato")][0]
        cells = [c.strip() for c in row.strip("|").split("|")]
        self.assertEqual(cells[4], "—")       # precision undefined, not 0
        self.assertEqual(cells[5], "0.000")   # recall genuinely zero
        self.assertEqual(refresh(out, self.tables), out)  # idempotent

    def test_check_fails_without_markers(self):
        import subprocess
        import sys
        from src.config import REPO_ROOT
        doc = Path(tempfile.mkdtemp()) / "no_tables.md"
        doc.write_text("# nothing here\n", encoding="utf-8")
        rel = doc.resolve()
        p = subprocess.run([sys.executable, "-m", "src.reporting.report_tables", "--check", "--report", str(rel)],
                           cwd=REPO_ROOT, capture_output=True, text=True)
        self.assertEqual(p.returncode, 1, p.stdout + p.stderr)


if __name__ == "__main__":
    unittest.main()
